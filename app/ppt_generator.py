import os
import json
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def generate_powerpoint_from_lesson(
    title: str,
    subject: str,
    grade: str,
    lesson_text: str,
    output_path: str,
    gemini_api_key: str = None
) -> dict:
    """
    Tạo file PowerPoint chuẩn sư phạm từ giáo án.
    Nếu có Gemini API Key, sử dụng AI phân tích cấu trúc sâu sắc.
    Nếu không, dùng bộ phân tích sư phạm tự động (Rule-Based NLP).
    """
    slides_data = []

    if gemini_api_key and gemini_api_key.strip():
        try:
            slides_data = _ai_generate_slides(title, subject, grade, lesson_text, gemini_api_key)
        except Exception as e:
            print(f"Gemini API warning, fallback to local rule-based: {e}")
            slides_data = _rule_based_generate_slides(title, subject, grade, lesson_text)
    else:
        slides_data = _rule_based_generate_slides(title, subject, grade, lesson_text)

    # Xây dựng file PPTX từ dữ liệu slide
    _build_pptx_file(title, subject, grade, slides_data, output_path)

    return {
        "title": title,
        "subject": subject,
        "grade": grade,
        "total_slides": len(slides_data),
        "slides": slides_data,
        "output_path": output_path
    }

def _ai_generate_slides(title: str, subject: str, grade: str, text: str, api_key: str) -> list:
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    prompt = f"""
Bạn là chuyên gia sư phạm thiết kế bài giảng điện tử. Hãy phân tích giáo án dưới đây và chuyển thành một chuỗi các slide PowerPoint hoàn chỉnh, sinh động, chuẩn cấu trúc sư phạm Việt Nam (Khởi động, Hình thành kiến thức, Luyện tập, Vận dụng).

Thông tin:
- Tên bài: {title}
- Môn học: {subject}
- Lớp: {grade}

Nội dung giáo án:
{text[:8000]}

Yêu cầu xuất ra định dạng JSON duy nhất với cấu trúc:
[
  {{
    "title": "Tên tiêu đề slide",
    "subtitle": "Phần phụ hoặc mô tả ngắn (nếu có)",
    "bullets": ["Ý chính 1", "Ý chính 2", "Ý chính 3", "Ý chính 4"],
    "key_takeaway": "Ghi nhớ quan trọng của slide này"
  }}
]
Chỉ trả về JSON thuần túy trong khối ```json ... ``` hoặc chuỗi JSON. Không thêm lời mở đầu hay kết thúc.
"""
    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=prompt
    )
    raw_text = response.text.strip()
    match = re.search(r"```json\s*(.*?)\s*```", raw_text, re.DOTALL)
    if match:
        raw_text = match.group(1)
    return json.loads(raw_text)

def _rule_based_generate_slides(title: str, subject: str, grade: str, text: str) -> list:
    """Phân tách giáo án thông minh không cần API key"""
    slides = []

    # Slide 1: Bìa
    slides.append({
        "title": f"BÀI HỌC: {title.upper()}",
        "subtitle": f"Môn: {subject} | Lớp: {grade}",
        "bullets": [
            "Chào mừng các em học sinh đến với tiết học",
            "Chuẩn bị đầy đủ sách vở và đồ dùng học tập",
            "Tập trung và tích cực phát biểu xây dựng bài"
        ],
        "key_takeaway": "Khởi đầu tiết học hào hứng và hiệu quả!"
    })

    # Phân tích nội dung theo dòng
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    current_section = "I. MỤC TIÊU BÀI HỌC"
    current_bullets = []

    for line in lines:
        if any(line.startswith(prefix) for prefix in ["I.", "II.", "III.", "IV.", "V.", "1.", "2.", "3.", "Phần ", "Hoạt động"]):
            if current_bullets:
                slides.append({
                    "title": current_section,
                    "subtitle": "Nội dung trọng tâm",
                    "bullets": current_bullets[:5],
                    "key_takeaway": "Nắm vững kiến thức cốt lõi"
                })
                current_bullets = []
            current_section = line
        else:
            if len(line) > 10 and len(current_bullets) < 5:
                current_bullets.append(line)

    if current_bullets:
        slides.append({
            "title": current_section,
            "subtitle": "Chi tiết bài học",
            "bullets": current_bullets[:5],
            "key_takeaway": "Ghi nhớ và vận dụng"
        })

    # Slide tổng kết & dặn dò nếu ít hơn 4 slide
    if len(slides) < 4:
        slides.append({
            "title": "LUYỆN TẬP & CỦNG CỐ",
            "subtitle": "Hệ thống hóa kiến thức",
            "bullets": [
                "Thảo luận và giải quyết các câu hỏi trọng tâm",
                "Phân tích ví dụ thực tế liên quan đến bài học",
                "Đánh giá và nhận xét kết quả hoạt động"
            ],
            "key_takeaway": "Khắc sâu kiến thức thông qua thực hành"
        })
        slides.append({
            "title": "HƯỚNG DẪN VỀ NHÀ",
            "subtitle": "Nhiệm vụ tự học",
            "bullets": [
                "Ôn lại toàn bộ lý thuyết đã học trong bài",
                "Hoàn thành các bài tập trong SGK và SBT",
                "Đọc trước nội dung bài học của tiết tiếp theo"
            ],
            "key_takeaway": "Tự học là chìa khóa của thành công"
        })

    return slides

def _build_pptx_file(title: str, subject: str, grade: str, slides_data: list, output_path: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 Widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Bảng màu hiện đại
    COLOR_PRIMARY = RGBColor(26, 86, 219)     # Deep Blue
    COLOR_ACCENT = RGBColor(14, 165, 233)     # Cyan/Sky
    COLOR_DARK = RGBColor(30, 41, 59)         # Slate 800
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_BG_CARD = RGBColor(248, 250, 252)   # Slate 50

    for idx, sdata in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)

        # Header background banner
        header_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.3)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = COLOR_PRIMARY
        header_shape.line.fill.background()

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.5), Inches(0.9))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = sdata.get("title", f"Slide {idx + 1}")
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.name = "Arial"

        # Subtitle badge if any
        sub = sdata.get("subtitle", "")
        if sub:
            p2 = tf.add_paragraph()
            p2.text = sub
            p2.font.size = Pt(14)
            p2.font.color.rgb = RGBColor(224, 231, 255)
            p2.font.name = "Arial"

        # Content Main Card
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.5)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = RGBColor(226, 232, 240)
        card.line.width = Pt(1.5)

        # Bullets inside Card
        content_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.0), Inches(3.9))
        ctf = content_box.text_frame
        ctf.word_wrap = True

        bullets = sdata.get("bullets", [])
        for b_idx, bullet in enumerate(bullets):
            bp = ctf.add_paragraph() if b_idx > 0 else ctf.paragraphs[0]
            bp.text = f"•  {bullet}"
            bp.font.size = Pt(20)
            bp.font.color.rgb = COLOR_DARK
            bp.font.name = "Arial"
            bp.space_after = Pt(14)

        # Key Takeaway Footer Banner if present
        takeaway = sdata.get("key_takeaway", "")
        if takeaway:
            tk_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.4), Inches(11.733), Inches(0.65)
            )
            tk_shape.fill.solid()
            tk_shape.fill.fore_color.rgb = RGBColor(238, 242, 255)
            tk_shape.line.color.rgb = COLOR_ACCENT
            tk_shape.line.width = Pt(1)

            tk_box = slide.shapes.add_textbox(Inches(1.1), Inches(6.42), Inches(11.0), Inches(0.6))
            tk_tf = tk_box.text_frame
            tk_p = tk_tf.paragraphs[0]
            tk_p.text = f"💡 Ghi nhớ: {takeaway}"
            tk_p.font.size = Pt(14)
            tk_p.font.bold = True
            tk_p.font.color.rgb = COLOR_PRIMARY
            tk_p.font.name = "Arial"

        # Slide Number
        num_box = slide.shapes.add_textbox(Inches(12.0), Inches(7.1), Inches(1.0), Inches(0.3))
        num_tf = num_box.text_frame
        num_p = num_tf.paragraphs[0]
        num_p.text = f"{idx + 1} / {len(slides_data)}"
        num_p.font.size = Pt(11)
        num_p.font.color.rgb = COLOR_MUTED
        num_p.alignment = PP_ALIGN.RIGHT

    prs.save(output_path)
